# -*- coding: utf-8 -*-
"""
Scraper de electrodomésticos · demo del Taller de IA (Secretaría de Juventud FUBA, 07/10/2026)

La pregunta: ¿el mismo producto cuesta lo mismo en todas las tiendas?

Qué hace, paso a paso:
  1. Busca un producto (por defecto, heladeras) en cuatro tiendas online argentinas.
  2. De cada resultado guarda: nombre, marca, precio, precio de lista, cuotas sin interés y link.
  3. Junta los que son EXACTAMENTE el mismo modelo: mismo código de barras (EAN).
  4. Calcula cuánto cambia el precio de una tienda a otra.
  5. Arma un Excel con la comparación y un resumen en pantalla.

Ojo con un detalle que apareció al mirar los datos: en estas webs no vende solo la tienda.
También venden otros comercios (como en un shopping). El scraper marca con * esos precios
y hace la comparación dos veces: con todo lo publicado y solo con lo que vende cada tienda.

Uso:
  python scraper_electrodomesticos.py                     heladeras
  python scraper_electrodomesticos.py lavarropas          cualquier otro producto
  python scraper_electrodomesticos.py "aire acondicionado"
  python scraper_electrodomesticos.py --offline           plan B: usa lo guardado en cache_electro/
  python scraper_electrodomesticos.py --no-abrir          no abre el Excel al terminar

Las cuatro tiendas usan la misma plataforma (VTEX), que publica el catálogo en un formato
ordenado (JSON). Sus robots.txt permiten leerlo. Igual vamos despacio: pocas páginas y una
pausa entre pedido y pedido.
"""
import datetime as dt
import json
import os
import sys
import time
import unicodedata
import urllib.parse
import urllib.request
from concurrent.futures import ThreadPoolExecutor

# ---------------------------------------------------------------------------
# Configuración
# ---------------------------------------------------------------------------
TIENDAS = {                        # nombre -> dirección
    'Frávega': 'www.fravega.com',
    'Naldo': 'www.naldo.com.ar',
    'Carrefour': 'www.carrefour.com.ar',
    'Pardo': 'www.pardo.com.ar',
}
PAGINAS = 2                        # 50 productos por página y por tienda
PAUSA = 0.4                        # segundos entre pedidos a una misma tienda
CARPETA = os.path.dirname(os.path.abspath(__file__))
CACHE = os.path.join(CARPETA, 'cache_electro')
NAVEGADOR = ('Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 '
             '(KHTML, like Gecko) Chrome/129.0 Safari/537.36')

ARGUMENTOS = [a for a in sys.argv[1:] if not a.startswith('--')]
BUSQUEDA = ARGUMENTOS[0] if ARGUMENTOS else 'heladera'
OFFLINE = '--offline' in sys.argv

os.system('')  # activa los colores en la consola de Windows
VERDE, ROJO, AMARILLO, GRIS, NEGRITA, FIN = '\033[92m', '\033[91m', '\033[93m', '\033[90m', '\033[1m', '\033[0m'


def pesos(x):
    """1234567.8 -> '$1.234.568'"""
    return '$' + format(round(x), ',').replace(',', '.')


def normalizar(texto):
    """Minúsculas y sin acentos, para comparar textos."""
    texto = unicodedata.normalize('NFD', texto.lower())
    return ''.join(c for c in texto if unicodedata.category(c) != 'Mn')


# ---------------------------------------------------------------------------
# 1. Bajar el catálogo (con copia local para el modo sin internet)
# ---------------------------------------------------------------------------
def descargar_pagina(tienda, numero):
    archivo = os.path.join(CACHE, '%s_%s_%d.json' % (normalizar(tienda), normalizar(BUSQUEDA).replace(' ', '-'), numero))
    if OFFLINE:
        if not os.path.exists(archivo):
            return []
        with open(archivo, encoding='utf-8') as f:
            return json.load(f)
    consulta = urllib.parse.urlencode({'ft': BUSQUEDA, '_from': numero * 50, '_to': numero * 50 + 49})
    url = 'https://%s/api/catalog_system/pub/products/search?%s' % (TIENDAS[tienda], consulta)
    pedido = urllib.request.Request(url, headers={'User-Agent': NAVEGADOR})
    with urllib.request.urlopen(pedido, timeout=30) as r:
        datos = json.loads(r.read().decode('utf-8'))
    os.makedirs(CACHE, exist_ok=True)
    with open(archivo, 'w', encoding='utf-8') as f:
        json.dump(datos, f, ensure_ascii=False)
    time.sleep(PAUSA)
    return datos


def descargar_tienda(tienda):
    """Todas las páginas de una tienda. Si la tienda falla, seguimos con las demás."""
    productos = []
    try:
        for numero in range(PAGINAS):
            pagina = descargar_pagina(tienda, numero)
            productos += pagina
            if len(pagina) < 50:
                break
    except Exception as error:
        print('  %s%s no respondió (%s). Sigo con las demás.%s' % (AMARILLO, tienda, error.__class__.__name__, FIN))
    return tienda, productos


# ---------------------------------------------------------------------------
# 2. Leer cada producto
# ---------------------------------------------------------------------------
def es_lo_que_busco(nombre):
    """Que sea una heladera y no 'Huevera para heladera': el nombre empieza con lo buscado."""
    palabras = normalizar(BUSQUEDA).split()
    nombre = normalizar(nombre)
    return nombre.startswith(palabras[0]) and all(p in nombre for p in palabras)


def cuotas_sin_interes(oferta):
    cuotas = [c['NumberOfInstallments'] for c in oferta.get('Installments', []) if c.get('InterestRate', 1) == 0]
    return max(cuotas) if cuotas else 1


def ofertas_de(tienda, producto):
    """Una oferta por cada variante (color, capacidad...) que tenga stock.
    Tomamos el precio que muestra la página: el del vendedor principal."""
    for item in producto.get('items', []):
        ean = (item.get('ean') or '').strip()
        vendedor = next((v for v in item.get('sellers', []) if v.get('sellerDefault')), None)
        if not ean or not vendedor:
            continue
        oferta = vendedor['commertialOffer']
        if oferta.get('AvailableQuantity', 0) <= 0 or oferta.get('Price', 0) <= 0:
            continue
        lista = max(oferta.get('ListPrice', 0), oferta.get('PriceWithoutDiscount', 0), oferta['Price'])
        yield {
            'tienda': tienda,
            'vende_la_tienda': vendedor.get('sellerId') == '1',     # en VTEX, el vendedor "1" es la propia tienda
            'producto': producto['productName'].strip(),
            'marca': producto.get('brand', ''),
            'ean': ean,
            'precio': oferta['Price'],
            'precio_lista': lista,
            'descuento': round((1 - oferta['Price'] / lista) * 100) if lista else 0,
            'cuotas': cuotas_sin_interes(oferta),
            'link': producto.get('link') or 'https://%s/%s/p' % (TIENDAS[tienda], producto.get('linkText', '')),
        }


# ---------------------------------------------------------------------------
# 3. Juntar el mismo modelo y comparar
# ---------------------------------------------------------------------------
def comparar(ofertas):
    por_codigo = {}
    for o in ofertas:
        actual = por_codigo.setdefault(o['ean'], {}).get(o['tienda'])
        if actual is None or o['precio'] < actual['precio']:
            por_codigo[o['ean']][o['tienda']] = o
    comparacion = []
    for ean, por_tienda in por_codigo.items():
        if len(por_tienda) < 2:
            continue
        orden = sorted(por_tienda.values(), key=lambda o: o['precio'])
        barata, cara = orden[0], orden[-1]
        comparacion.append({
            'ean': ean,
            'producto': barata['producto'],
            'marca': barata['marca'],
            'precios': {t: o['precio'] for t, o in por_tienda.items()},
            'links': {t: o['link'] for t, o in por_tienda.items()},
            'de_la_tienda': {t: o['vende_la_tienda'] for t, o in por_tienda.items()},
            'mas_barata': barata['tienda'],
            'mas_cara': cara['tienda'],
            'diferencia': round((cara['precio'] / barata['precio'] - 1) * 100, 1),
            'ahorro': cara['precio'] - barata['precio'],
        })
    comparacion.sort(key=lambda c: c['diferencia'], reverse=True)
    return comparacion


# ---------------------------------------------------------------------------
# 4. Guardar: Excel y un resumen en JSON
# ---------------------------------------------------------------------------
def guardar_excel(ofertas, comparacion, solo_tienda, ruta, hoy):
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.styles import Alignment, Font, PatternFill

    amarillo = PatternFill('solid', fgColor='EBB416')
    verde, rojo = PatternFill('solid', fgColor='C6EFCE'), PatternFill('solid', fgColor='F4CCCC')

    def encabezado(hoja, columnas):
        hoja.append([c for c, _ in columnas])
        for i, (_, ancho) in enumerate(columnas, start=1):
            celda = hoja.cell(1, i)
            celda.font, celda.fill = Font(bold=True), amarillo
            celda.alignment = Alignment(wrap_text=True, vertical='center')
            hoja.column_dimensions[celda.column_letter].width = ancho
        hoja.freeze_panes = 'A2'

    libro = Workbook()

    tiendas = list(TIENDAS)

    def hoja_comparacion(hoja, filas):
        encabezado(hoja, [('Producto', 50), ('Marca', 13)] + [(t, 14) for t in tiendas] +
                   [('Más barata', 12), ('Más cara', 12), ('Diferencia %', 12), ('Te ahorrás', 14), ('Código de barras', 16)])
        for c in filas:
            hoja.append([c['producto'], c['marca']] + [c['precios'].get(t) for t in tiendas] +
                        [c['mas_barata'], c['mas_cara'], c['diferencia'] / 100, c['ahorro'], c['ean']])
            fila = hoja.max_row
            for j, t in enumerate(tiendas, start=3):
                celda = hoja.cell(fila, j)
                if t in c['precios']:
                    # con * si lo vende otro comercio dentro de la web de la tienda
                    celda.number_format = '"$"#,##0' if c['de_la_tienda'][t] else '"$"#,##0" *"'
                    celda.hyperlink = c['links'][t]          # clic en el precio -> página del producto
                    celda.fill = verde if t == c['mas_barata'] else (rojo if t == c['mas_cara'] else PatternFill())
            hoja.cell(fila, 3 + len(tiendas) + 2).number_format = '0.0%'
            hoja.cell(fila, 3 + len(tiendas) + 3).number_format = '"$"#,##0'
        hoja.auto_filter.ref = hoja.dimensions
        hoja.cell(hoja.max_row + 2, 1).value = '* Lo vende otro comercio dentro de la web de la tienda (marketplace), no la tienda.'

    # Hoja 1: el mismo modelo en varias tiendas, con todo lo publicado
    hoja = libro.active
    hoja.title = 'Mismo modelo'
    hoja_comparacion(hoja, comparacion)

    # Gráfico: las 10 diferencias más grandes
    if comparacion:
        top = min(10, len(comparacion))
        grafico = BarChart()
        grafico.type = 'bar'
        grafico.title = 'Mismo modelo: cuánto más cuesta en la tienda más cara'
        grafico.y_axis.number_format = '0%'
        grafico.legend = None
        col = 3 + len(tiendas) + 2
        grafico.add_data(Reference(hoja, min_col=col, min_row=1, max_row=top + 1), titles_from_data=True)
        grafico.set_categories(Reference(hoja, min_col=1, min_row=2, max_row=top + 1))
        grafico.height, grafico.width = 9, 22
        hoja.add_chart(grafico, 'A%d' % (hoja.max_row + 3))

    # Hoja 2: solo lo que vende cada tienda
    hoja_comparacion(libro.create_sheet('Solo la tienda'), solo_tienda)

    # Hoja 3: todo lo que encontró
    todos = libro.create_sheet('Todos los productos')
    encabezado(todos, [('Tienda', 12), ('Vendido por', 14), ('Producto', 55), ('Marca', 13), ('Precio', 14),
                       ('Precio de lista', 15), ('Descuento %', 12), ('Cuotas sin interés', 12), ('Código de barras', 16), ('Link', 45)])
    for o in sorted(ofertas, key=lambda o: (o['tienda'], o['precio'])):
        todos.append([o['tienda'], 'La tienda' if o['vende_la_tienda'] else 'Otro comercio', o['producto'], o['marca'],
                      o['precio'], o['precio_lista'], o['descuento'] / 100, o['cuotas'], o['ean'], o['link']])
        fila = todos.max_row
        todos.cell(fila, 5).number_format = todos.cell(fila, 6).number_format = '"$"#,##0'
        todos.cell(fila, 7).number_format = '0%'
        todos.cell(fila, 10).hyperlink = o['link']
    todos.auto_filter.ref = todos.dimensions

    # Hoja 4: quién fue la más barata (solo con lo que vende cada tienda)
    quien = libro.create_sheet('Quién fue más barata')
    encabezado(quien, [('Tienda', 14), ('Modelos que vende y comparte con otra', 22), ('Fue la más barata', 16), ('Fue la más cara', 16)])
    for t in tiendas:
        tiene = sum(1 for c in solo_tienda if t in c['precios'])
        quien.append([t, tiene, sum(1 for c in solo_tienda if c['mas_barata'] == t), sum(1 for c in solo_tienda if c['mas_cara'] == t)])

    nota = libro.create_sheet('Fuente')
    nota.column_dimensions['A'].width = 110
    for linea in ['Precios publicados en %s el %s.' % (', '.join(TIENDAS.values()), hoy.strftime('%d/%m/%Y')),
                  'Es el precio que muestra cada página para el vendedor principal, con stock. No incluye envío ni promociones bancarias.',
                  '"Mismo modelo" = mismo código de barras (EAN).',
                  '* = lo vende otro comercio dentro de la web de la tienda (marketplace). La hoja "Solo la tienda" compara solo lo que vende cada tienda.',
                  'Búsqueda: "%s".' % BUSQUEDA]:
        nota.append([linea])
    libro.save(ruta)


# ---------------------------------------------------------------------------
# Todo junto
# ---------------------------------------------------------------------------
def main():
    hoy = dt.date.today()
    print('\n%sBuscando "%s" en %d tiendas%s  (%s)%s\n' % (NEGRITA, BUSQUEDA, len(TIENDAS), FIN, hoy.strftime('%d/%m/%Y'),
                                                       '  [modo sin internet]' if OFFLINE else ''))

    # Paso 1: las cuatro tiendas a la vez
    with ThreadPoolExecutor(max_workers=len(TIENDAS)) as grupo:
        catalogos = list(grupo.map(descargar_tienda, TIENDAS))

    # Paso 2: quedarnos con lo que buscamos, con stock y código de barras
    ofertas = []
    for tienda, productos in catalogos:
        propias = [o for p in productos if es_lo_que_busco(p.get('productName', '')) for o in ofertas_de(tienda, p)]
        ofertas += propias
        de_la_tienda = sum(1 for o in propias if o['vende_la_tienda'])
        print('  %-10s %3d publicaciones con precio y stock  (%d las vende la tienda; %d, otros comercios)'
              % (tienda, len(propias), de_la_tienda, len(propias) - de_la_tienda))

    # Paso 3: el mismo modelo en varias tiendas
    comparacion = comparar(ofertas)
    solo_tienda = comparar([o for o in ofertas if o['vende_la_tienda']])
    modelos = len({o['ean'] for o in ofertas})
    print('\n  %d modelos distintos. %s%d están en más de una tienda%s: los comparo.\n' % (modelos, NEGRITA, len(comparacion), FIN))

    def marca(c, t):
        return '' if c['de_la_tienda'][t] else '*'

    print('  %sMismo modelo, distinto precio%s   (* = lo vende otro comercio en esa web)' % (NEGRITA, FIN))
    for c in comparacion[:10]:
        b, k = c['mas_barata'], c['mas_cara']
        print('  %s+%5.1f%%%s  %-46s %s%s %s%s%s  vs  %s%s %s%s%s' % (
            ROJO, c['diferencia'], FIN, c['producto'][:46],
            VERDE, b, pesos(c['precios'][b]), marca(c, b), FIN, GRIS, k, pesos(c['precios'][k]), marca(c, k), FIN))

    # Paso 4: guardar
    nombre = 'electro_%s_%s' % (normalizar(BUSQUEDA).replace(' ', '-'), hoy.isoformat())
    excel = os.path.join(CARPETA, nombre + '.xlsx')
    guardar_excel(ofertas, comparacion, solo_tienda, excel, hoy)
    resumen = {'busqueda': BUSQUEDA, 'fecha': hoy.isoformat(), 'tiendas': list(TIENDAS),
               'productos': len(ofertas), 'modelos': modelos, 'repetidos': len(comparacion),
               'comparacion': comparacion, 'solo_tienda': solo_tienda}
    with open(os.path.join(CARPETA, nombre + '.json'), 'w', encoding='utf-8') as f:
        json.dump(resumen, f, ensure_ascii=False, indent=1)

    def resumir(titulo, filas):
        if not filas:
            return
        promedio = sum(c['diferencia'] for c in filas) / len(filas)
        mayor = filas[0]
        print('  %s%s%s' % (NEGRITA, titulo, FIN))
        print('    %d modelos en más de una tienda · diferencia promedio %s%.0f%%%s · máxima %s%.0f%%%s'
              % (len(filas), NEGRITA, promedio, FIN, ROJO, mayor['diferencia'], FIN))
        print('    La mayor: %s (%s %s vs. %s %s)' % (mayor['producto'][:44], mayor['mas_barata'],
              pesos(mayor['precios'][mayor['mas_barata']]), mayor['mas_cara'], pesos(mayor['precios'][mayor['mas_cara']])))

    print('\n%sResumen%s  ·  %d publicaciones en %d tiendas\n' % (NEGRITA, FIN, len(ofertas), len(TIENDAS)))
    resumir('Con todo lo publicado en cada web', comparacion)
    resumir('Solo con lo que vende cada tienda', solo_tienda)
    print('\n  Planilla: %s\n' % os.path.basename(excel))

    if '--no-abrir' not in sys.argv and hasattr(os, 'startfile'):
        os.startfile(excel)


if __name__ == '__main__':
    main()
